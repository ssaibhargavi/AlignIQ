import pdfplumber
import docx

def extract_from_pdf(pdf_file) -> str:
    """Extracts text from PDF files using pdfplumber."""
    text_content = []
    with pdfplumber.open(pdf_file) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
    return "\n".join(text_content).strip()

def extract_from_docx(docx_file) -> str:
    """
    Extracts text from Word documents (.docx).
    Also extracts text from tables (since many Word resumes use tables for skills/education).
    """
    doc = docx.Document(docx_file)
    text_content = []

    # 1. Extract regular paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            text_content.append(p.text.strip())

    # 2. Extract table content (if any)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                # Remove duplicate cell text if cells were merged
                unique_cells = list(dict.fromkeys(row_texts))
                text_content.append(" | ".join(unique_cells))

    return "\n".join(text_content).strip()

def extract_from_txt(txt_file) -> str:
    """Extracts text from plain text (.txt) files."""
    try:
        # Read raw bytes and decode with fallback
        return txt_file.read().decode("utf-8", errors="ignore").strip()
    except Exception:
        txt_file.seek(0)
        return txt_file.read().decode("latin-1", errors="ignore").strip()

def extract_text_from_file(uploaded_file) -> str:
    """
    Universal multi-format text extractor.
    Accepts PDF, DOCX, or TXT and returns clean extracted text.
    """
    if uploaded_file is None:
        return ""

    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_from_pdf(uploaded_file)
    elif file_name.endswith(".docx"):
        return extract_from_docx(uploaded_file)
    elif file_name.endswith(".txt"):
        return extract_from_txt(uploaded_file)
    else:
        return ""
